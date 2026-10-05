import io
from datetime import datetime
from functools import partial

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

# Columnas disponibles para armar un listado de productos, en el mismo
# orden en que se ofrecen en la app. La clave es lo que manda/recibe la
# API; el valor es la etiqueta legible que se usa como encabezado.
COLUMNAS_LISTADO = {
    "nombre": "Nombre",
    "categoria_nombre": "Categoría",
    "marca_nombre": "Marca",
    "proveedor_nombre": "Proveedor",
    "precio_unitario": "Precio unitario",
    "porcentaje_ganancia": "% Ganancia",
    "precio_venta": "Precio venta",
    "stock": "Stock",
}

_COLUMNAS_DECIMALES = {"precio_unitario", "porcentaje_ganancia", "precio_venta"}


def _valor_columna(producto, campo):
    if campo == "categoria_nombre":
        return producto.categoria.nombre if producto.categoria_id else ""
    if campo == "marca_nombre":
        return producto.marca.nombre if producto.marca_id else ""
    if campo == "proveedor_nombre":
        return producto.proveedor.nombre if producto.proveedor_id else ""
    if campo in _COLUMNAS_DECIMALES:
        return float(getattr(producto, campo))
    return getattr(producto, campo)


def _texto_pdf(campo, valor):
    if campo in _COLUMNAS_DECIMALES:
        return f"{valor:.2f}"
    return str(valor)


def generar_listado_excel(productos, columnas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Listado"

    ws.append([COLUMNAS_LISTADO[c] for c in columnas])
    for celda in ws[1]:
        celda.font = Font(bold=True)

    for producto in productos:
        ws.append([_valor_columna(producto, c) for c in columnas])

    for idx, campo in enumerate(columnas, start=1):
        letra = ws.cell(row=1, column=idx).column_letter
        if campo in _COLUMNAS_DECIMALES:
            for fila in range(2, ws.max_row + 1):
                ws.cell(row=fila, column=idx).number_format = "0.00"
        longitudes = [
            len(str(ws.cell(row=f, column=idx).value))
            for f in range(1, ws.max_row + 1)
            if ws.cell(row=f, column=idx).value is not None
        ]
        ws.column_dimensions[letra].width = min((max(longitudes) if longitudes else 10) + 4, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"listado_productos_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    respuesta = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre_archivo}"'
    return respuesta


_TAMANIO_PAGINA_PDF = landscape(A4)


def _dibujar_pie_de_pagina(canvas, doc, fecha_hora):
    """
    Pie de página chico y discreto, igual en todas las hojas: fecha/hora
    de generación abajo a la izquierda (sin ningún texto como "Generado
    el", solo el dato) y número de página (solo el número, sin "Página"
    ni total) abajo a la derecha.
    """
    ancho_pagina, _ = _TAMANIO_PAGINA_PDF
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#999999"))
    canvas.drawString(1.5 * cm, 0.8 * cm, fecha_hora)
    canvas.drawRightString(ancho_pagina - 1.5 * cm, 0.8 * cm, str(canvas.getPageNumber()))
    canvas.restoreState()


def generar_listado_pdf(productos, columnas):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=_TAMANIO_PAGINA_PDF,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph("Listado de productos", estilos["Title"]),
        HRFlowable(
            width="100%",
            thickness=1,
            color=colors.HexColor("#2E3A46"),
            spaceBefore=2,
            spaceAfter=8,
        ),
        Spacer(1, 0.3 * cm),
    ]

    filas = [[COLUMNAS_LISTADO[c] for c in columnas]]
    for producto in productos:
        filas.append([_texto_pdf(c, _valor_columna(producto, c)) for c in columnas])

    tabla = Table(filas, repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E3A46")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    elementos.append(tabla)

    fecha_hora = datetime.now().strftime("%d/%m/%Y %H:%M")
    pie_de_pagina = partial(_dibujar_pie_de_pagina, fecha_hora=fecha_hora)
    doc.build(elementos, onFirstPage=pie_de_pagina, onLaterPages=pie_de_pagina)
    buffer.seek(0)

    nombre_archivo = f"listado_productos_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
    respuesta = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre_archivo}"'
    return respuesta
